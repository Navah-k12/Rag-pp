import os, shutil, json
from fastapi import FastAPI, HTTPException, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from dotenv import load_dotenv
from google import genai

load_dotenv()

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
UPLOAD_DIR = os.path.join(BASE_DIR, "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)

TEXT_PATH = os.path.join(UPLOAD_DIR, "documento_transcrito.txt")
META_PATH = os.path.join(UPLOAD_DIR, "metadata.json")
CHUNKS_PATH = os.path.join(UPLOAD_DIR, "chunks.json")

IMAGE_EXTS = {'.png', '.jpg', '.jpeg', '.webp', '.gif', '.bmp', '.tiff'}
DOC_EXTS   = {'.pdf', '.pptx', '.txt'}

app = FastAPI(title="StudyRAG Backend API")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

api_key = os.getenv("GEMINI_API_KEY")
client = None
if api_key and api_key != "tu_api_key_aqui":
    client = genai.Client(api_key=api_key)
else:
    print("⚠️ GEMINI_API_KEY no configurada.")

class QuestionRequest(BaseModel):
    question: str

class FlashcardRequest(BaseModel):
    count: int = 5

class QuizRequest(BaseModel):
    count: int = 5
    type: str = "multiple_choice"  # multiple_choice | true_false | mixed

# ─── helpers ─────────────────────────────────────────────────

def load_text():
    if not os.path.exists(TEXT_PATH):
        raise HTTPException(status_code=400, detail="Primero sube un documento con POST /upload")
    with open(TEXT_PATH, "r", encoding="utf-8") as f:
        return f.read()

def save_metadata(filename, filetype):
    meta = {"filename": filename, "filetype": filetype}
    with open(META_PATH, "w", encoding="utf-8") as f:
        json.dump(meta, f)

def extract_text(filepath, ext):
    if ext in IMAGE_EXTS:
        from PIL import Image
        imagen = Image.open(filepath)
        prompt = (
            "Transcribe de forma exacta todo el texto contenido en esta imagen. "
            "Mantén el formato original, listas, títulos y estructura. "
            "No agregues explicaciones tuyas, solo el texto transcrito directamente."
        )
        response = client.models.generate_content(
            model='gemini-2.0-flash',
            contents=[imagen, prompt]
        )
        return response.text
    elif ext == '.pdf':
        import pdfplumber
        text = ""
        with pdfplumber.open(filepath) as pdf:
            for i, page in enumerate(pdf.pages, 1):
                t = page.extract_text() or ""
                text += f"\n--- Página {i} ---\n{t}"
        return text.strip()
    elif ext == '.pptx':
        from pptx import Presentation
        prs = Presentation(filepath)
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text)
            parts.append(f"--- Diapositiva {i} ---\n" + "\n".join(texts))
        return "\n\n".join(parts)
    elif ext == '.txt':
        with open(filepath, "r", encoding="utf-8") as f:
            return f.read()
    return ""

def chunk_text(text, size=500, overlap=50):
    words = text.split()
    chunks = []
    i = 0
    while i < len(words):
        chunk = " ".join(words[i:i + size])
        chunks.append({"index": len(chunks), "text": chunk})
        i += size - overlap
    return chunks

def transcribir_documento(filepath, filename):
    global client
    if not client:
        raise HTTPException(status_code=500, detail="Gemini API Key no configurada.")

    ext = os.path.splitext(filename)[1].lower()
    print(f"📄 Procesando {filename}...")
    try:
        text = extract_text(filepath, ext)
        if not text.strip():
            raise ValueError("No se pudo extraer texto del documento.")

        chunks = chunk_text(text)
        with open(TEXT_PATH, "w", encoding="utf-8") as f:
            f.write(text)
        with open(CHUNKS_PATH, "w", encoding="utf-8") as f:
            json.dump(chunks, f)
        save_metadata(filename, ext)
        print(f"✅ Documento procesado: {len(chunks)} chunks, {len(text)} caracteres")
        return {"message": "Documento procesado con éxito", "chunks": len(chunks), "chars": len(text)}
    except Exception as e:
        print(f"❌ Error: {e}")
        raise

# ─── endpoints ───────────────────────────────────────────────

@app.get("/")
def read_root():
    return {"message": "Servidor StudyRAG en línea"}

@app.get("/status")
def get_status():
    api_ready = client is not None
    transcribed = os.path.exists(TEXT_PATH)
    meta = {"filename": None, "filetype": None}
    if os.path.exists(META_PATH):
        with open(META_PATH) as f:
            meta = json.load(f)
    preview = ""
    if transcribed:
        with open(TEXT_PATH, "r", encoding="utf-8") as f:
            preview = f.read(500)
    return {
        "api_key_configured": api_ready,
        "document_transcribed": transcribed,
        "document": meta,
        "preview": preview if transcribed else "Aún no has subido ningún documento"
    }

@app.post("/upload")
async def upload_file(file: UploadFile = File(...)):
    if not client:
        raise HTTPException(status_code=500, detail="Gemini API Key no configurada.")

    ext = os.path.splitext(file.filename)[1].lower()
    allowed = IMAGE_EXTS | DOC_EXTS
    if ext not in allowed:
        raise HTTPException(status_code=400,
            detail=f"Formato no soportado. Permitidos: {', '.join(allowed)}")

    filepath = os.path.join(UPLOAD_DIR, f"upload{ext}")
    with open(filepath, "wb") as f:
        shutil.copyfileobj(file.file, f)

    return transcribir_documento(filepath, file.filename)

@app.post("/ask")
def ask_question(req: QuestionRequest):
    if not client:
        raise HTTPException(status_code=500, detail="Gemini API Key no configurada.")

    document_text = load_text()
    meta = {}
    if os.path.exists(META_PATH):
        with open(META_PATH) as f:
            meta = json.load(f)

    prompt = f"""Eres un asistente de estudio útil. Responde la pregunta del usuario basándote ÚNICAMENTE en el documento provisto.

Si la respuesta no está en el documento, di: "La respuesta a esta pregunta no se encuentra en el documento proporcionado."
Cuando cites información, indica la página o sección de donde proviene.

--- INICIO DEL DOCUMENTO ({meta.get('filename', 'documento')}) ---
{document_text}
--- FIN DEL DOCUMENTO ---

Pregunta: {req.question}
Respuesta:"""

    try:
        response = client.models.generate_content(
            model='gemini-2.0-flash',
            contents=prompt
        )
        return {"answer": response.text}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error al consultar el modelo: {e}")

@app.post("/summarize")
def summarize():
    if not client:
        raise HTTPException(status_code=500, detail="Gemini API Key no configurada.")
    document_text = load_text()

    prompt = f"""Genera un resumen estructurado del siguiente documento. Incluye:
1. Ideas principales
2. Puntos clave por sección/tema
3. Conclusión

--- DOCUMENTO ---
{document_text[:10000]}
--- FIN ---

Resumen:"""

    try:
        response = client.models.generate_content(model='gemini-2.0-flash', contents=prompt)
        return {"summary": response.text}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error: {e}")

@app.post("/flashcards")
def generate_flashcards(req: FlashcardRequest):
    if not client:
        raise HTTPException(status_code=500, detail="Gemini API Key no configurada.")
    document_text = load_text()

    prompt = f"""Genera {req.count} flashcards de estudio basadas en el siguiente documento.
Responde ÚNICAMENTE con un array JSON. Cada flashcard debe tener los campos:
- "front": la pregunta o concepto
- "back": la respuesta o definición

--- DOCUMENTO ---
{document_text[:10000]}
--- FIN ---

JSON:"""

    try:
        response = client.models.generate_content(model='gemini-2.0-flash', contents=prompt)
        text = response.text.strip()
        if text.startswith("```json"): text = text[7:]
        if text.startswith("```"):     text = text[3:]
        if text.endswith("```"):       text = text[:-3]
        cards = json.loads(text.strip())
        return {"flashcards": cards}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error al generar flashcards: {e}")

@app.post("/quiz")
def generate_quiz(req: QuizRequest):
    if not client:
        raise HTTPException(status_code=500, detail="Gemini API Key no configurada.")
    document_text = load_text()

    t_inst = ""
    if req.type == "multiple_choice":
        t_inst = "Todas las preguntas deben ser de opción múltiple con 4 opciones (A, B, C, D)."
    elif req.type == "true_false":
        t_inst = "Todas las preguntas deben ser de verdadero/falso."
    else:
        t_inst = "Mezcla preguntas de opción múltiple y verdadero/falso."

    prompt = f"""Genera un examen de {req.count} preguntas basado en el siguiente documento.
{t_inst}
Responde ÚNICAMENTE con un array JSON. Cada pregunta debe tener:
- "question": el texto de la pregunta
- "options": array de opciones (para true_false: ["Verdadero", "Falso"])
- "answer": la respuesta correcta
- "explanation": breve explicación

--- DOCUMENTO ---
{document_text[:10000]}
--- FIN ---

JSON:"""

    try:
        response = client.models.generate_content(model='gemini-2.0-flash', contents=prompt)
        text = response.text.strip()
        if text.startswith("```json"): text = text[7:]
        if text.startswith("```"):     text = text[3:]
        if text.endswith("```"):       text = text[:-3]
        questions = json.loads(text.strip())
        return {"quiz": questions}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error al generar quiz: {e}")
