import os

from app.core.config import IMAGE_EXTS


def extract_text(filepath: str, ext: str) -> str:
    if ext in IMAGE_EXTS:
        from PIL import Image

        from app.services.rag_pipeline import generate

        imagen = Image.open(filepath)
        prompt = (
            "Transcribe de forma exacta todo el texto contenido en esta imagen. "
            "Mantén el formato original, listas, títulos y estructura. "
            "No agregues explicaciones tuyas, solo el texto transcrito directamente."
        )
        return generate(prompt, image=imagen)

    if ext == ".pdf":
        return _extract_pdf(filepath)

    if ext == ".pptx":
        return _extract_pptx(filepath)

    if ext == ".txt":
        with open(filepath, "r", encoding="utf-8") as f:
            return f.read()

    return ""


def _extract_pdf(filepath: str) -> str:
    import pdfplumber

    from app.services.rag_pipeline import generate, gemini_client

    text = ""
    with pdfplumber.open(filepath) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            t = page.extract_text() or ""
            text += f"\n--- Página {i} ---\n{t}"
    text = text.strip()

    if len(text) < 100 and gemini_client:
        print("⚠️ pdfplumber no extrajo texto, usando Gemini OCR para el PDF...")
        return _ocr_pdf(filepath)

    return text


def _ocr_pdf(filepath: str) -> str:
    from PIL import Image
    import fitz

    from app.services.rag_pipeline import generate

    doc = fitz.open(filepath)
    ocr_text = ""
    for i, page in enumerate(doc, 1):
        pix = page.get_pixmap(dpi=150)
        img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        prompt = (
            f"Transcribe de forma exacta todo el texto de la página {i} de {len(doc)} del PDF. "
            "Mantén el formato original. Solo el texto, sin explicaciones."
        )
        page_text = generate(prompt, image=img)
        ocr_text += f"\n--- Página {i} ---\n{page_text}"
    doc.close()
    return ocr_text.strip()


def _extract_pptx(filepath: str) -> str:
    from pptx import Presentation

    prs = Presentation(filepath)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text)
        parts.append(f"--- Diapositiva {i} ---\n" + "\n".join(texts))
    return "\n\n".join(parts)


def chunk_text(text: str, size: int = 500, overlap: int = 50) -> list[dict]:
    words = text.split()
    chunks = []
    i = 0
    while i < len(words):
        chunk = " ".join(words[i : i + size])
        chunks.append({"index": len(chunks), "text": chunk})
        i += size - overlap
    return chunks
